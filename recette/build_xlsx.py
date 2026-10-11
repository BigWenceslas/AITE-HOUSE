# -*- coding: utf-8 -*-
"""Génère le cahier de recette AITE HOUSE (.xlsx)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from cas_de_test import CAS, DOMAINES

TEAL = "0F6478"
TEAL_D = "0A4C5C"
PEACH = "F4CDAA"
VIOLET = "7A2E86"
GRIS = "EEF2F4"
BLANC = "FFFFFF"

F = "Arial"
STATUTS = ["À tester", "OK", "KO", "Bloqué", "Non applicable"]
PRIO_LBL = {1: "1 - Critique", 2: "2 - Majeur", 3: "3 - Mineur"}
DOM_NOM = {c: n for c, n, _ in DOMAINES}
DOM_MOD = {c: m for c, _, m in DOMAINES}

thin = Side(style="thin", color="C9D4DA")
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ═══════════════ 1. SYNTHÈSE ═══════════════
ws = wb.active
ws.title = "Synthèse"
ws.sheet_view.showGridLines = False

ws["B2"] = "CAHIER DE RECETTE — AITE HOUSE"
ws["B2"].font = Font(name=F, size=18, bold=True, color=TEAL_D)
ws["B3"] = "Suite de gestion hôtelière sur Odoo 18 — 12 modules"
ws["B3"].font = Font(name=F, size=11, color="5A6670")
ws["B4"] = "AITE Consulting SARL · info@aite-consulting.com · (+237) 673 373 367"
ws["B4"].font = Font(name=F, size=9, color="8A949C")

# bandeau d'identification (cellules à remplir = fond jaune)
ident = [("Client / établissement", ""), ("Version testée", "AITE HOUSE — septembre 2026"),
         ("Environnement", "Recette"), ("Base de données", ""),
         ("Date de début", ""), ("Date de fin", ""),
         ("Testeur principal", ""), ("Responsable AITE", "")]
ws["B6"] = "IDENTIFICATION DE LA RECETTE"
ws["B6"].font = Font(name=F, size=10, bold=True, color=BLANC)
ws["B6"].fill = PatternFill("solid", fgColor=TEAL)
ws.merge_cells("B6:E6")
r = 7
for lbl, val in ident:
    ws.cell(r, 2, lbl).font = Font(name=F, size=10, bold=True)
    c = ws.cell(r, 3, val)
    c.font = Font(name=F, size=10)
    c.fill = PatternFill("solid", fgColor="FFF9D6")
    c.border = BORD
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    r += 1

ws["B16"] = "Cellules sur fond jaune : à compléter. Saisir les résultats dans l'onglet « Cas de test »."
ws["B16"].font = Font(name=F, size=9, italic=True, color="8A949C")

# tableau d'avancement global
ws["B18"] = "AVANCEMENT GLOBAL"
ws["B18"].font = Font(name=F, size=10, bold=True, color=BLANC)
ws["B18"].fill = PatternFill("solid", fgColor=TEAL)
ws.merge_cells("B18:E18")
ws["B19"] = "Statut"
ws["C19"] = "Nombre"
ws["D19"] = "Part"
for col in "BCD":
    c = ws[f"{col}19"]
    c.font = Font(name=F, size=10, bold=True, color=TEAL_D)
    c.fill = PatternFill("solid", fgColor=GRIS)
    c.border = BORD
row = 20
for s in STATUTS:
    ws.cell(row, 2, s).font = Font(name=F, size=10)
    ws.cell(row, 3).value = f"=COUNTIF('Cas de test'!$I$2:$I$200,$B{row})"
    ws.cell(row, 4).value = f"=IFERROR($C{row}/$C$25,0)"
    ws.cell(row, 4).number_format = "0.0%"
    for col in (2, 3, 4):
        ws.cell(row, col).border = BORD
    row += 1
ws.cell(25, 2, "TOTAL").font = Font(name=F, size=10, bold=True)
ws.cell(25, 3).value = "=COUNTA('Cas de test'!$A$2:$A$200)"
ws.cell(25, 3).font = Font(name=F, size=10, bold=True)
ws.cell(25, 4).value = "=IFERROR(SUM($D$20:$D$24),0)"
ws.cell(25, 4).number_format = "0.0%"
ws.cell(25, 4).font = Font(name=F, size=10, bold=True)
for col in (2, 3, 4):
    ws.cell(25, col).border = BORD
    ws.cell(25, col).fill = PatternFill("solid", fgColor=PEACH)

ws["B27"] = "Taux de réussite (OK / testés)"
ws["B27"].font = Font(name=F, size=10, bold=True)
ws["C27"] = "=IFERROR($C$21/($C$21+$C$22),0)"
ws["C27"].number_format = "0.0%"
ws["C27"].font = Font(name=F, size=12, bold=True, color=TEAL_D)
ws["B28"] = "Cas critiques en échec (bloquants)"
ws["B28"].font = Font(name=F, size=10, bold=True)
ws["C28"] = ("=COUNTIFS('Cas de test'!$D$2:$D$200,\"1 - Critique\","
             "'Cas de test'!$I$2:$I$200,\"KO\")")
ws["C28"].font = Font(name=F, size=12, bold=True, color="C0392B")

# avancement par domaine
ws["G18"] = "AVANCEMENT PAR DOMAINE"
ws["G18"].font = Font(name=F, size=10, bold=True, color=BLANC)
ws["G18"].fill = PatternFill("solid", fgColor=TEAL)
ws.merge_cells("G18:L18")
heads = ["Domaine", "Module", "Cas", "OK", "KO", "Reste à tester"]
for i, h in enumerate(heads):
    c = ws.cell(19, 7 + i, h)
    c.font = Font(name=F, size=10, bold=True, color=TEAL_D)
    c.fill = PatternFill("solid", fgColor=GRIS)
    c.border = BORD
row = 20
for code, nom, mod in DOMAINES:
    ws.cell(row, 7, nom).font = Font(name=F, size=9)
    ws.cell(row, 8, mod).font = Font(name=F, size=9, color="5A6670")
    ws.cell(row, 9).value = f"=COUNTIF('Cas de test'!$B$2:$B$200,$G{row})"
    ws.cell(row, 10).value = (
        f"=COUNTIFS('Cas de test'!$B$2:$B$200,$G{row},"
        f"'Cas de test'!$I$2:$I$200,\"OK\")")
    ws.cell(row, 11).value = (
        f"=COUNTIFS('Cas de test'!$B$2:$B$200,$G{row},"
        f"'Cas de test'!$I$2:$I$200,\"KO\")")
    ws.cell(row, 12).value = (
        f"=COUNTIFS('Cas de test'!$B$2:$B$200,$G{row},"
        f"'Cas de test'!$I$2:$I$200,\"À tester\")")
    for col in range(7, 13):
        ws.cell(row, col).border = BORD
        ws.cell(row, col).font = Font(name=F, size=9)
    row += 1

for col, w in {"A": 2, "B": 34, "C": 22, "D": 12, "E": 10, "F": 3,
               "G": 40, "H": 26, "I": 8, "J": 8, "K": 8, "L": 14}.items():
    ws.column_dimensions[col].width = w

# ═══════════════ 2. CAS DE TEST ═══════════════
ws2 = wb.create_sheet("Cas de test")
ws2.sheet_view.showGridLines = False
COLS = [("ID", 8), ("Domaine", 30), ("Module", 24), ("Priorité", 12),
        ("Pré-requis", 34), ("Scénario", 38), ("Étapes", 52),
        ("Résultat attendu", 52), ("Statut", 14), ("Testé par", 14),
        ("Date", 12), ("Anomalie / commentaire", 34)]
for i, (h, w) in enumerate(COLS, start=1):
    c = ws2.cell(1, i, h)
    c.font = Font(name=F, size=10, bold=True, color=BLANC)
    c.fill = PatternFill("solid", fgColor=TEAL_D)
    c.alignment = Alignment(horizontal="center", vertical="center",
                            wrap_text=True)
    c.border = BORD
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.row_dimensions[1].height = 28

for r, (cid, dom, prio, pre, scen, etapes, attendu) in enumerate(CAS,
                                                                start=2):
    vals = [cid, DOM_NOM[dom], DOM_MOD[dom], PRIO_LBL[prio], pre, scen,
            etapes, attendu, "À tester", "", "", ""]
    for i, v in enumerate(vals, start=1):
        c = ws2.cell(r, i, v)
        c.font = Font(name=F, size=9)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = BORD
        if i == 1:
            c.font = Font(name=F, size=9, bold=True, color=TEAL_D)
            c.alignment = Alignment(horizontal="center", vertical="top")
        if i == 4 and prio == 1:
            c.font = Font(name=F, size=9, bold=True, color="C0392B")
        if i == 6:
            c.font = Font(name=F, size=9, bold=True)
        if i in (9, 10, 11, 12):
            c.fill = PatternFill("solid", fgColor="FFF9D6")
    ws2.row_dimensions[r].height = 58

dv = DataValidation(type="list", formula1='"%s"' % ",".join(STATUTS),
                    allow_blank=True, showDropDown=False)
dv.error = "Choisir un statut dans la liste."
ws2.add_data_validation(dv)
dv.add(f"I2:I{len(CAS) + 1}")
ws2.freeze_panes = "B2"
ws2.auto_filter.ref = f"A1:L{len(CAS) + 1}"

# ═══════════════ 3. ANOMALIES ═══════════════
ws3 = wb.create_sheet("Anomalies")
ws3.sheet_view.showGridLines = False
ACOLS = [("N°", 7), ("Cas lié", 10), ("Module", 24), ("Description", 52),
         ("Gravité", 14), ("Détecté le", 12), ("Détecté par", 16),
         ("Statut", 16), ("Corrigé le", 12), ("Vérifié par", 16)]
for i, (h, w) in enumerate(ACOLS, start=1):
    c = ws3.cell(1, i, h)
    c.font = Font(name=F, size=10, bold=True, color=BLANC)
    c.fill = PatternFill("solid", fgColor=VIOLET)
    c.alignment = Alignment(horizontal="center", vertical="center",
                            wrap_text=True)
    c.border = BORD
    ws3.column_dimensions[get_column_letter(i)].width = w
ws3.row_dimensions[1].height = 26

# ligne d'exemple (format attendu)
exemple = [1, "E03", "aite_hotel_pos_link",
           "Exemple : une chambre libre apparaît dans la liste des notes "
           "de chambre.", "Bloquante", "12/09/2026", "Nom du testeur",
           "Ouverte", "", ""]
for i, v in enumerate(exemple, start=1):
    c = ws3.cell(2, i, v)
    c.font = Font(name=F, size=9, italic=True, color="8A949C")
    c.alignment = Alignment(vertical="top", wrap_text=True)
    c.border = BORD
ws3.row_dimensions[2].height = 34
ws3.cell(3, 1, "↑ ligne d'exemple à remplacer — saisir les anomalies à "
               "partir de la ligne 3").font = Font(name=F, size=9,
                                                   italic=True,
                                                   color="8A949C")
dv2 = DataValidation(type="list",
                     formula1='"Bloquante,Majeure,Mineure,Cosmétique"',
                     allow_blank=True, showDropDown=False)
ws3.add_data_validation(dv2)
dv2.add("E2:E200")
dv3 = DataValidation(type="list",
                     formula1='"Ouverte,En cours,Corrigée,Vérifiée,'
                              'Rejetée"',
                     allow_blank=True, showDropDown=False)
ws3.add_data_validation(dv3)
dv3.add("H2:H200")
ws3.freeze_panes = "A2"
ws3.auto_filter.ref = "A1:J200"

# ═══════════════ 4. MODE D'EMPLOI ═══════════════
ws4 = wb.create_sheet("Mode d'emploi")
ws4.sheet_view.showGridLines = False
ws4.column_dimensions["B"].width = 100
lignes = [
    ("COMMENT UTILISER CE CAHIER DE RECETTE", "t"),
    ("", ""),
    ("1. Préparer l'environnement", "h"),
    ("Installer la suite AITE HOUSE sur une base de recette (jamais sur la "
     "production). Installer aite_demo_data pour disposer de données, ou "
     "partir des données réelles du client selon le contexte.", ""),
    ("", ""),
    ("2. Dérouler les cas dans l'ordre des domaines", "h"),
    ("Les cas A à M suivent le parcours naturel de l'établissement : "
     "hébergement, espaces, réservation en ligne, caisse, note de chambre, "
     "crédit, fidélité, stock, achats, pilotage, puis transverse.", ""),
    ("Traiter en priorité les cas « 1 - Critique » : ils conditionnent la "
     "mise en production.", ""),
    ("", ""),
    ("3. Saisir le résultat de chaque cas", "h"),
    ("Colonne Statut (liste déroulante) : À tester / OK / KO / Bloqué / "
     "Non applicable.", ""),
    ("Renseigner le testeur, la date, et en cas d'échec une description "
     "précise dans la colonne Anomalie.", ""),
    ("", ""),
    ("4. Journaliser les anomalies", "h"),
    ("Chaque KO donne lieu à une ligne dans l'onglet Anomalies, avec sa "
     "gravité et son statut de correction. La ligne 2 est un exemple à "
     "remplacer.", ""),
    ("", ""),
    ("5. Suivre l'avancement", "h"),
    ("L'onglet Synthèse se met à jour automatiquement : avancement global, "
     "taux de réussite, cas critiques en échec, et détail par domaine.", ""),
    ("", ""),
    ("6. Critères de prononcé de la recette", "h"),
    ("· Recette prononcée : 100 % des cas critiques en OK et aucune "
     "anomalie bloquante ouverte.", ""),
    ("· Recette prononcée avec réserves : cas critiques en OK, anomalies "
     "majeures ou mineures listées avec une date de correction convenue.",
     ""),
    ("· Recette refusée : au moins un cas critique en KO.", ""),
    ("", ""),
    ("Éditeur : AITE Consulting SARL — Douala, Makepe Bloc F — "
     "info@aite-consulting.com", "f"),
]
r = 2
for txt, typ in lignes:
    c = ws4.cell(r, 2, txt)
    if typ == "t":
        c.font = Font(name=F, size=15, bold=True, color=TEAL_D)
    elif typ == "h":
        c.font = Font(name=F, size=11, bold=True, color=TEAL)
    elif typ == "f":
        c.font = Font(name=F, size=9, italic=True, color="8A949C")
    else:
        c.font = Font(name=F, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws4.row_dimensions[r].height = 30
    r += 1

# mise en page d'impression : paysage, ajusté à la largeur
for sheet in (ws, ws2, ws3, ws4):
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
ws2.print_title_rows = "1:1"
ws3.print_title_rows = "1:1"

OUT_DIR = os.environ.get("RECETTE_OUT", os.path.dirname(os.path.abspath(__file__)))
wb.save(os.path.join(OUT_DIR, "AITE_HOUSE_cahier_de_recette.xlsx"))
print("classeur écrit — %d cas" % len(CAS))
